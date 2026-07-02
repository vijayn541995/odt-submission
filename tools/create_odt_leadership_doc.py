from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/vn105957/Desktop/odt-submission")
OUT_DIR = ROOT / "docs"
ASSET_DIR = OUT_DIR / "generated-assets"
DOCX_PATH = OUT_DIR / "ODT-2.0-Leadership-Brief.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1F2933"
MUTED = "5B6770"
LINE = "D9E2EC"
FILL = "F2F4F7"
SOFT_BLUE = "E8EEF5"
GREEN = "1F7A4D"
RED = "C74634"
GOLD = "A06B00"
TEAL = "0F6B5C"


def ensure_dirs() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)


def font(size: int, bold: bool = False) -> ImageFont.ImageFont:
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Supplemental/Helvetica Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Helvetica.ttf",
        "/Library/Fonts/Arial Bold.ttf" if bold else "/Library/Fonts/Arial.ttf",
    ]
    for item in candidates:
        try:
            return ImageFont.truetype(item, size)
        except Exception:
            continue
    return ImageFont.load_default()


def draw_round_rect(draw: ImageDraw.ImageDraw, xy, radius, fill, outline, width=2):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def text_center(draw: ImageDraw.ImageDraw, xy, text, font_obj, fill=INK, line_gap=6):
    x1, y1, x2, y2 = xy
    lines = text.split("\n")
    heights = []
    widths = []
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font_obj)
        widths.append(bbox[2] - bbox[0])
        heights.append(bbox[3] - bbox[1])
    total_h = sum(heights) + line_gap * (len(lines) - 1)
    y = y1 + ((y2 - y1) - total_h) / 2
    for line, w, h in zip(lines, widths, heights):
        draw.text((x1 + ((x2 - x1) - w) / 2, y), line, font=font_obj, fill=fill)
        y += h + line_gap


def arrow(draw: ImageDraw.ImageDraw, start, end, fill="#6B7280", width=3):
    draw.line([start, end], fill=fill, width=width)
    sx, sy = start
    ex, ey = end
    if abs(ex - sx) >= abs(ey - sy):
        direction = 1 if ex > sx else -1
        pts = [(ex, ey), (ex - 12 * direction, ey - 7), (ex - 12 * direction, ey + 7)]
    else:
        direction = 1 if ey > sy else -1
        pts = [(ex, ey), (ex - 7, ey - 12 * direction), (ex + 7, ey - 12 * direction)]
    draw.polygon(pts, fill=fill)


def architecture_diagram(path: Path) -> None:
    img = Image.new("RGB", (1600, 950), "white")
    d = ImageDraw.Draw(img)
    title_f = font(38, True)
    head_f = font(24, True)
    body_f = font(21)
    small_f = font(18)

    d.text((60, 40), "ODT 2.0 Architecture: Governed SDLC Control Plane", font=title_f, fill=f"#{DARK_BLUE}")
    d.text((60, 92), "ODT orchestrates workflow, evidence, policy, and agent execution while worker engines perform scoped tasks.", font=body_f, fill=f"#{MUTED}")

    boxes = {
        "ui": (80, 190, 400, 360, "Developer Workbench\nReact + Vite UI", f"#{SOFT_BLUE}"),
        "api": (500, 190, 850, 360, "ODT Backend API\nNode 24 + local services", "#FFF7ED"),
        "db": (955, 190, 1280, 360, "Evidence Store\nSQLite / future DB", "#EEF8F6"),
        "policy": (500, 455, 850, 625, "Governance Engine\nStandards, gates, approvals", "#FEF3C7"),
        "agents": (955, 455, 1280, 625, "Agent Orchestrator\nCodex now, Cline/OCI later", "#FEE2E2"),
        "connectors": (80, 455, 400, 625, "Connector Gateway\nJira, Git, Build, MCP", "#EFF6FF"),
        "models": (955, 720, 1280, 880, "AI Providers\nLocal, OCI GenAI, Ollama,\nOpenAI-compatible", "#F5F3FF"),
        "workers": (1320, 455, 1540, 625, "Workers\nPlanner\nFull Stack Dev\nReviewer\nBuild Verifier", "#F8FAFC"),
    }

    for key, (x1, y1, x2, y2, label, fill) in boxes.items():
        draw_round_rect(d, (x1, y1, x2, y2), 20, fill, f"#{LINE}", 3)
        first, *rest = label.split("\n")
        d.text((x1 + 24, y1 + 24), first, font=head_f, fill=f"#{DARK_BLUE}")
        d.multiline_text((x1 + 24, y1 + 66), "\n".join(rest), font=body_f, fill=f"#{INK}", spacing=8)

    arrow(d, (400, 275), (500, 275))
    arrow(d, (850, 275), (955, 275))
    arrow(d, (675, 360), (675, 455))
    arrow(d, (850, 540), (955, 540))
    arrow(d, (1280, 540), (1320, 540))
    arrow(d, (240, 455), (240, 360))
    arrow(d, (1120, 625), (1120, 720))
    arrow(d, (955, 800), (850, 600))

    d.text((60, 845), "Control principle: analyze locally first, route minimal scoped context to the right worker, record evidence and decisions.", font=small_f, fill=f"#{MUTED}")
    img.save(path)


def workflow_diagram(path: Path) -> None:
    img = Image.new("RGB", (1600, 900), "white")
    d = ImageDraw.Draw(img)
    title_f = font(38, True)
    head_f = font(22, True)
    small_f = font(18)

    d.text((60, 40), "ODT Workflow: Requirement to PR Readiness", font=title_f, fill=f"#{DARK_BLUE}")
    d.text((60, 92), "Each stage produces evidence. Gates decide whether the work can continue, delegate, rework, or stop.", font=small_f, fill=f"#{MUTED}")

    stages = [
        ("1", "Intake", "Requirement, Jira,\nrepo, assets"),
        ("2", "Analyze", "Repo evidence,\ngaps, context"),
        ("3", "Plan", "Design,\nimplementation plan"),
        ("4", "Standards Gate", "Security, tests,\naccessibility, policy"),
        ("5", "Human Approval", "Write scope,\nrisk acceptance"),
        ("6", "Delegate", "Codex worker\nhandoff"),
        ("7", "Ingest", "Logs, output,\nquestions"),
        ("8", "Review/Rework", "Reviewer findings,\nrelay context"),
        ("9", "Build Verify", "Approved tests,\nbuild proof"),
        ("10", "PR Ready", "Evidence-backed\nPR package"),
    ]

    x0, y0 = 70, 190
    box_w, box_h = 260, 130
    gap_x, gap_y = 45, 95
    coords = []
    for i, (num, title, desc) in enumerate(stages):
        row = 0 if i < 5 else 1
        col = i if i < 5 else 9 - i
        x = x0 + col * (box_w + gap_x)
        y = y0 + row * (box_h + gap_y)
        coords.append((x, y, x + box_w, y + box_h))
        fill = "#FFF7ED" if title in {"Standards Gate", "Human Approval"} else f"#{SOFT_BLUE}"
        if title in {"Review/Rework"}:
            fill = "#FEF3C7"
        if title in {"PR Ready"}:
            fill = "#EEF8F6"
        draw_round_rect(d, (x, y, x + box_w, y + box_h), 18, fill, f"#{LINE}", 3)
        d.ellipse((x + 18, y + 18, x + 58, y + 58), fill=f"#{RED}" if title in {"Standards Gate", "Human Approval"} else f"#{DARK_BLUE}")
        text_center(d, (x + 18, y + 18, x + 58, y + 58), num, font(18, True), "white")
        d.text((x + 72, y + 20), title, font=head_f, fill=f"#{DARK_BLUE}")
        d.multiline_text((x + 72, y + 58), desc, font=small_f, fill=f"#{INK}", spacing=6)

    for i in range(4):
        arrow(d, (coords[i][2], (coords[i][1] + coords[i][3]) // 2), (coords[i + 1][0], (coords[i + 1][1] + coords[i + 1][3]) // 2))
    arrow(d, ((coords[4][0] + coords[4][2]) // 2, coords[4][3]), ((coords[5][0] + coords[5][2]) // 2, coords[5][1]))
    for i in range(5, 9):
        arrow(d, (coords[i][0], (coords[i][1] + coords[i][3]) // 2), (coords[i + 1][2], (coords[i + 1][1] + coords[i + 1][3]) // 2))

    # Rework loop.
    d.arc((635, 505, 965, 790), start=15, end=335, fill=f"#{GOLD}", width=5)
    arrow(d, (665, 620), (620, 585), fill=f"#{GOLD}", width=4)
    d.text((650, 765), "If review finds gaps, ODT sends scoped relay context back to the right worker.", font=small_f, fill=f"#{GOLD}")

    d.text((60, 835), "Current implementation focus: completed-Jira verification uses Reviewer first, then Build Verifier/test evidence, then PR Ready.", font=small_f, fill=f"#{MUTED}")
    img.save(path)


def cover_banner(path: Path) -> None:
    img = Image.new("RGB", (1600, 520), "#FFFDFB")
    d = ImageDraw.Draw(img)
    title_f = font(48, True)
    body_f = font(24)
    small_f = font(19)

    # Subtle product-style bands.
    d.rectangle((0, 0, 1600, 520), fill="#FFFDFB")
    d.rounded_rectangle((52, 56, 1548, 464), radius=34, fill="#F7F3EF", outline=f"#{LINE}", width=3)
    d.rectangle((52, 56, 88, 464), fill=f"#{RED}")
    d.rounded_rectangle((92, 96, 252, 256), radius=34, fill=f"#{RED}")
    d.text((127, 132), "ODT", font=font(42, True), fill="white")
    for cx in (136, 172, 208):
        d.ellipse((cx - 8, 218 - 8, cx + 8, 218 + 8), fill="white")

    d.text((310, 102), "Oracle Developer Twin 2.0", font=title_f, fill=f"#{DARK_BLUE}")
    d.text((310, 166), "Governed agentic SDLC workbench for enterprise development teams", font=body_f, fill=f"#{MUTED}")

    card_specs = [
        (310, 260, 620, 392, "Governed", "Standards, approval, risk,\nand audit evidence"),
        (650, 260, 960, 392, "Agentic", "Codex worker delegation\nwith relay context"),
        (990, 260, 1300, 392, "Efficient", "Local evidence + scoped\nprompts reduce token waste"),
    ]
    for x1, y1, x2, y2, head, copy in card_specs:
        d.rounded_rectangle((x1, y1, x2, y2), radius=18, fill="white", outline=f"#{LINE}", width=2)
        d.text((x1 + 22, y1 + 22), head, font=font(24, True), fill=f"#{DARK_BLUE}")
        d.multiline_text((x1 + 22, y1 + 58), copy, font=small_f, fill=f"#{INK}", spacing=6)

    d.text((310, 420), "Control plane first. Worker engines second. Evidence always.", font=small_f, fill=f"#{TEAL}")
    img.save(path)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=90, start=130, bottom=90, end=130) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        margin = tc_mar.find(qn(f"w:{margin_name}"))
        if margin is None:
            margin = OxmlElement(f"w:{margin_name}")
            tc_mar.append(margin)
        margin.set(qn("w:w"), str(value))
        margin.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = LINE, size: str = "8") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def repeat_header_row(table) -> None:
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def set_table_width(table, width_dxa: int = 9360, indent_dxa: int = 120) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False


def format_cell_text(cell, *, size: float = 9.5, bold: bool = False, color: str = INK) -> None:
    set_cell_margins(cell)
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.05
        for run in paragraph.runs:
            run.font.name = "Calibri"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = RGBColor.from_string(color)


def style_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_title(doc: Document, banner_path: Path) -> None:
    doc.add_picture(str(banner_path), width=Inches(6.5))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("ODT 2.0: Governed Agentic SDLC Workbench")
    r.font.name = "Calibri"
    r.font.size = Pt(24)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(DARK_BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("Leadership and Engineering Team Brief")
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor.from_string(MUTED)

    meta = doc.add_table(rows=2, cols=2)
    set_table_width(meta)
    set_table_borders(meta)
    values = [
        ("Prepared for", "Engineering leadership and development teams"),
        ("Date", "June 10, 2026"),
        ("Status", "ODT 2.0 local workbench in active evolution"),
        ("Primary message", "Agentic development with governance, auditability, and policy awareness"),
    ]
    for idx, (k, v) in enumerate(values):
        row = idx // 2
        col = idx % 2
        cell = meta.cell(row, col)
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        label = p.add_run(k.upper())
        label.font.size = Pt(7.5)
        label.font.bold = True
        label.font.color.rgb = RGBColor.from_string(MUTED)
        p = cell.add_paragraph()
        value = p.add_run(v)
        value.font.size = Pt(10)
        value.font.bold = True if k in {"Status", "Primary message"} else False
        value.font.color.rgb = RGBColor.from_string(DARK_BLUE if k in {"Status", "Primary message"} else INK)
        set_cell_shading(cell, "F8FAFC")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, top=120, start=150, bottom=120, end=150)
    for row in meta.rows:
        set_cell_width(row.cells[0], 4680)
        set_cell_width(row.cells[1], 4680)

    add_value_cards(doc)


def add_value_cards(doc: Document) -> None:
    doc.add_paragraph()
    table = doc.add_table(rows=1, cols=3)
    set_table_width(table)
    set_table_borders(table)
    values = [
        ("Problem", "AI coding is powerful, but requirements, standards, approvals, evidence, and PR readiness remain fragmented."),
        ("ODT Answer", "ODT acts as the governed SDLC control plane and delegates scoped work to Codex and future workers."),
        ("Expected Value", "Faster planning, lower repeated prompting, stronger quality gates, and better auditability for AI-assisted work."),
    ]
    fills = ["FEE2E2", "E8EEF5", "EEF8F6"]
    widths = [3120, 3120, 3120]
    for i, (head, copy) in enumerate(values):
        cell = table.cell(0, i)
        set_cell_width(cell, widths[i])
        set_cell_shading(cell, fills[i])
        set_cell_margins(cell, top=130, start=150, bottom=130, end=150)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(head)
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(DARK_BLUE)
        p = cell.add_paragraph(copy)
        p.paragraph_format.space_after = Pt(0)


def add_callout(doc: Document, title: str, body: str, fill: str = "EEF8F6") -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table)
    set_table_borders(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=130, start=150, bottom=130, end=150)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    p = cell.add_paragraph(body)
    p.paragraph_format.space_after = Pt(0)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_width(table)
    set_table_borders(table)
    repeat_header_row(table)
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr[i].text = header
        set_cell_shading(hdr[i], FILL)
        set_cell_width(hdr[i], widths[i])
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        format_cell_text(hdr[i], size=9.5, bold=True, color=DARK_BLUE)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            set_cell_width(cells[i], widths[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            format_cell_text(cells[i], size=9.3, color=INK)


def add_footer(doc: Document) -> None:
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.text = "ODT 2.0 Leadership Brief"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string(MUTED)
    footer = section.footer.paragraphs[0]
    footer.text = "Confidential internal draft - validate pilot metrics before external claims"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)


def build_doc() -> None:
    ensure_dirs()
    banner_path = ASSET_DIR / "odt-cover-banner.png"
    arch_path = ASSET_DIR / "odt-architecture.png"
    workflow_path = ASSET_DIR / "odt-workflow.png"
    cover_banner(banner_path)
    architecture_diagram(arch_path)
    workflow_diagram(workflow_path)

    doc = Document()
    style_doc(doc)
    add_footer(doc)
    add_title(doc, banner_path)

    doc.add_heading("1. Executive Summary", level=1)
    doc.add_paragraph(
        "Oracle Developer Twin 2.0, or ODT 2.0, is a governed agentic SDLC workbench that helps development teams move from requirement intake to PR readiness with traceability, policy gates, human approval, and controlled AI delegation."
    )
    doc.add_paragraph(
        "ODT does not replace Codex, Cline, OCI GenAI, or similar engines. It makes them usable as supervised workers inside a structured development workflow."
    )
    add_callout(
        doc,
        "Positioning",
        "ODT is the SDLC control plane. Codex, Cline, OCI GenAI, Ollama, OpenAI-compatible models, and future tools are worker engines that ODT can route, govern, and audit."
    )

    doc.add_heading("2. Problem ODT Solves", level=1)
    doc.add_paragraph("Enterprise development teams lose time and quality when SDLC work is fragmented across Jira, code repositories, documentation, standards, chat, AI sessions, build tools, and PR systems.")
    add_bullets(doc, [
        "Requirements are often incomplete, ambiguous, or separated from implementation context.",
        "Developers repeatedly assemble Jira, repo, standards, test, and review context by hand.",
        "AI coding tools often receive large prompts with repeated context, increasing token cost and model confusion.",
        "Standards such as security, dependency policy, accessibility, VPAT, WCAG, and testing are inconsistently applied.",
        "Review and PR readiness decisions are often subjective and poorly evidenced.",
        "Teams lack a durable audit trail for what AI did, what humans approved, and why the workflow continued."
    ])

    doc.add_heading("3. How ODT Helps Development Teams", level=1)
    add_bullets(doc, [
        "Creates a structured requirement-to-PR workflow instead of relying on ad hoc prompts.",
        "Captures requirement, Jira, repository, uploaded asset, standards, worker, review, and PR evidence.",
        "Runs governance gates before write delegation and PR readiness.",
        "Delegates scoped work to Codex worker lanes with handoff bundles and controlled context.",
        "Captures worker output, logs, questions, findings, and relay context for follow-on workers.",
        "Supports reviewer comments, rework routing, accepted risk, and build/test verification.",
        "Reduces repeated prompt construction by compiling only relevant context for each workflow stage."
    ])

    doc.add_heading("4. Expected Productivity and Quality Impact", level=1)
    doc.add_paragraph("The following are pilot targets, not guaranteed production claims. They should be validated through controlled team trials.")
    add_table(
        doc,
        ["Area", "Expected improvement", "How ODT contributes"],
        [
            ["Requirement understanding", "Fewer missed assumptions and late clarifications", "Structured intake, gap detection, clarification gates, Jira/repo context."],
            ["Planning effort", "20-40% reduction in repeated manual planning effort", "Reusable design, implementation plan, risk, standards, and test artifacts."],
            ["Prompt preparation", "30-60% reduction in repeated context assembly", "Local evidence store, scoped worker prompts, and relay context."],
            ["Review readiness", "Faster reviewer onboarding", "Evidence-backed review queue, worker output, standards findings, and PR checklist."],
            ["PR preparation", "25-50% faster PR summary/evidence preparation", "PR Ready package assembled from actual evidence."],
            ["Quality and compliance", "Improved consistency", "Policy gates for security, accessibility, dependency, testing, and maintainability."]
        ],
        [2200, 2500, 4660],
    )

    doc.add_heading("5. Current ODT Capabilities", level=1)
    add_bullets(doc, [
        "ODT Workbench UI with workflow navigation, Overview, Intake, Planner, Standards, Agent Team, Review, PR Ready, Artifacts, Guide, Runs, Monitoring, and Settings.",
        "Jira2 connector detection and Jira issue read path.",
        "Completed-Jira verification workflow for work that is already marked Done and has repository evidence.",
        "Structured planning, technical design, and standards evidence.",
        "Standards gate with accessibility, security, dependency, testing, maintainability, approval, and PR readiness concerns.",
        "Agent Team lanes including Lead Planner, Senior Full Stack Dev, Reviewer, and Build Verifier.",
        "Codex worker delegation wiring, handoff bundles, visible worker launch, refresh, stop, ingest, and evidence extraction flow.",
        "Relay context between workers so questions/findings can move from one lane to another.",
        "Review comments, blockers, accepted-risk evidence, and rework routing.",
        "PR Ready page with evidence checklist and PR package builder.",
        "Monitoring health log and backend-triggered UI smoke validation.",
        "Local-first evidence database and local deterministic guide behavior when no remote AI provider is configured."
    ])

    doc.add_page_break()
    doc.add_heading("6. Architecture", level=1)
    doc.add_paragraph("ODT is designed as a control plane around agent engines, evidence sources, and governance policy.")
    doc.add_picture(str(arch_path), width=Inches(6.5))

    add_table(
        doc,
        ["Layer", "Current implementation", "Planned evolution"],
        [
            ["Frontend workbench", "React 18, Vite, CSS, local browser UI.", "Role-aware views, richer dashboards, enterprise auth-aware UI."],
            ["Backend API", "Node.js 24, built-in node:sqlite, local HTTP API.", "Service deployment, durable DB, RBAC, multi-project tenancy."],
            ["Evidence store", "Local SQLite database and file artifacts.", "Shared enterprise database, object storage, audit export."],
            ["Agent orchestration", "Codex worker launch and handoff bundles.", "Cline, OCI GenAI/OCA, Ollama, OpenAI-compatible worker adapters."],
            ["Connectors", "Jira2 readiness/read path and connector health model.", "GitHub/Bitbucket, build service, knowledge, MCP-native connectors."],
            ["Validation", "Backend UI smoke validation and monitoring health log.", "Broader regression pack, worker-flow validation, policy-pack tests."]
        ],
        [2100, 3600, 3660],
    )

    doc.add_page_break()
    doc.add_heading("7. Workflow", level=1)
    doc.add_paragraph("ODT drives work through explicit stages. Each stage creates evidence, and gates decide whether the next action is allowed.")
    doc.add_picture(str(workflow_path), width=Inches(6.5))
    add_numbered(doc, [
        "Intake captures requirement, Jira, repo, and supporting assets.",
        "Analyze collects repository and work-state evidence.",
        "Plan drafts design, implementation tasks, risks, test strategy, and open questions.",
        "Standards Gate checks team/company policy expectations.",
        "Human Approval records write approval, risk acceptance, or stop decisions.",
        "Delegate sends scoped work to Codex or another worker engine.",
        "Ingest captures worker logs, output, questions, findings, and evidence.",
        "Review/Rework routes findings back to the correct worker lane when needed.",
        "Build Verify captures approved test/build proof.",
        "PR Ready packages evidence only when readiness is justified."
    ])

    doc.add_page_break()
    doc.add_heading("8. How ODT Differs From Codex, Cline, and Similar Tools", level=1)
    add_table(
        doc,
        ["Capability", "Codex / Cline / Cursor", "ODT 2.0"],
        [
            ["Primary role", "Worker engine for coding, commands, and technical assistance.", "Governed SDLC control plane."],
            ["Requirement handling", "Prompt/session dependent.", "Structured intake, gaps, assumptions, and clarification gates."],
            ["Governance", "Mostly manual or prompt-based.", "First-class standards gates, approvals, and risk decisions."],
            ["Agent orchestration", "Possible but ad hoc.", "Worker lanes, relay context, and evidence capture."],
            ["PR readiness", "Generated text or manual judgment.", "Evidence-backed readiness gate."],
            ["Token usage", "Large repeated prompts likely.", "Scoped context compiler and prompt budget roadmap."],
            ["Auditability", "Chat/log history.", "Structured evidence, run history, approvals, standards, and PR artifacts."]
        ],
        [2200, 3300, 3860],
    )

    doc.add_heading("9. Customization and Policy Packs", level=1)
    doc.add_paragraph("A major ODT differentiator is team and company customization. ODT can evolve into a policy-pack driven platform where company, organization, team, repo, and task-specific rules are inherited and overridden intentionally.")
    add_bullets(doc, [
        "Coding standards and approved patterns.",
        "Security and dependency approval rules.",
        "Accessibility, WCAG, VPAT, and Section 508 gates.",
        "Testing coverage and build verification requirements.",
        "Do/don't guidance for each team or repo.",
        "Allowed commands, restricted paths, and dependency install policy.",
        "Reviewer roles, approval ownership, and escalation rules.",
        "Prompt templates and worker instructions per domain."
    ])

    doc.add_page_break()
    doc.add_heading("10. Privacy, Token Cost, and Context Efficiency", level=1)
    add_callout(
        doc,
        "Context compiler principle",
        "ODT should analyze locally first, store structured evidence, and send only the minimal relevant context to each model or worker."
    )
    add_bullets(doc, [
        "Repo analysis, Jira evidence, standards checks, review comments, and worker outputs can remain local or in approved internal infrastructure.",
        "Codex and future model prompts can be scoped to the worker role instead of sending the entire conversation and repository context.",
        "Repeated prompt assembly is reduced because ODT reuses saved evidence.",
        "Token usage can be logged by worker, provider, model, workflow stage, and assignment.",
        "Privacy risk is reduced because raw context does not need to be manually pasted into generic AI sessions."
    ])

    doc.add_heading("11. Scalability Considerations", level=1)
    add_table(
        doc,
        ["Dimension", "Current local design", "Scalable direction"],
        [
            ["Projects", "Single local workbench database.", "Multi-project workspaces with project isolation and shared policy packs."],
            ["Users", "Local developer use.", "Authentication, authorization, roles, and approval ownership."],
            ["Evidence", "SQLite and local artifacts.", "Durable DB plus object storage for artifacts, logs, prompts, and reports."],
            ["Agents", "Codex worker launch wired.", "Adapter layer for Codex, Cline, OCI GenAI, Ollama, OpenAI-compatible models."],
            ["Connectors", "Jira2 configured and monitored.", "MCP-native Jira/Git/build/knowledge connectors with read/write gates."],
            ["Policy", "Baseline standards checks.", "Versioned company/team/repo policy registry with inheritance."],
            ["Observability", "Monitoring page and UI smoke runs.", "Central metrics, health dashboards, audit exports, and cost reporting."]
        ],
        [2300, 3300, 3760],
    )

    doc.add_page_break()
    doc.add_heading("12. Roadmap", level=1)
    add_table(
        doc,
        ["Phase", "Focus", "Expected outcome"],
        [
            ["Near term", "Stabilize Reviewer and Build Verifier closeout, PR evidence gaps, Jira connector reliability, and UI smoke coverage.", "Demo-grade ODT becomes a more reliable team pilot."],
            ["Mid term", "Policy packs, stronger ODT Guide RAG, Git/build connectors, Cline adapter, context compression, token budgeting.", "Teams can customize ODT and use multiple workers safely."],
            ["Long term", "Enterprise auth, RBAC, durable shared storage, MCP-native orchestration, OCI GenAI/OCA integration, audit exports, executive metrics.", "ODT becomes an enterprise-ready governed agentic SDLC platform."]
        ],
        [1600, 4400, 3360],
    )

    doc.add_heading("13. What Teams Can Expect", level=1)
    add_bullets(doc, [
        "Faster movement from requirement to implementation plan.",
        "Reduced repeated context gathering and prompt writing.",
        "Better visibility into what AI workers did and what evidence supports the work.",
        "More consistent standards compliance across teams.",
        "Earlier clarification of missing requirements and risks.",
        "Safer Codex/Cline usage through ODT approval gates and scoped handoff bundles.",
        "More reliable PR readiness through evidence-backed checks."
    ])

    doc.add_heading("14. Recommended Pilot Metrics", level=1)
    add_table(
        doc,
        ["Metric", "Why it matters"],
        [
            ["Time from Jira intake to implementation plan", "Measures planning acceleration."],
            ["Number of clarification gaps caught before coding", "Measures requirement quality improvement."],
            ["Prompt size per worker run", "Measures token and cost efficiency."],
            ["Review comments requiring rework", "Measures quality and standards effectiveness."],
            ["PR readiness blocker count", "Measures evidence and compliance completeness."],
            ["Manual context assembly time", "Measures developer productivity gain."],
            ["Accepted-risk decisions with notes", "Measures auditability and governance maturity."]
        ],
        [3300, 6060],
    )

    doc.add_heading("15. Closing Message", level=1)
    doc.add_paragraph(
        "ODT 2.0 is designed to make modern AI development practical for enterprise teams. It does not compete with Codex or Cline as a coding engine; it turns those engines into governed, auditable, policy-aware workers inside a disciplined SDLC flow."
    )
    doc.add_paragraph(
        "The leadership framing should be simple: ODT helps teams adopt agentic AI without losing control of standards, evidence, approval, privacy, cost, or PR readiness."
    )
    add_callout(
        doc,
        "Recommended next move",
        "Approve a controlled ODT 2.0 pilot with two or three development teams. Measure requirement clarity, manual context assembly time, review rework, PR readiness blockers, token usage, and audit completeness before scaling.",
        fill="FFF7ED",
    )

    doc.save(DOCX_PATH)


if __name__ == "__main__":
    build_doc()
    print(DOCX_PATH)
